"""Text-only writer (BLD-01, BLD-02, BLD-08): what the writer puts on disk."""

import re
import zipfile
from xml.etree import ElementTree

import openpyxl
import pytest
from epb.formats import probematrix_v2 as fmt
from epb.formats import validate as validate_module
from epb.formats import writer as writer_module
from epb.formats.validate import Finding, validate_workbook
from epb.formats.writer import PackInvalid, write_pack
from epbtest.packkit import CREATED, metadata, sheet_xml, valid_row, with_values

MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
EDGE_VALUES = ("=SUM(1,2)", "+44", "-5", "@x", "0", "40000000000000000001", "2026-09")


def _write(tmp_path, rows, name="pack.xlsx", **meta):
    path = tmp_path / name
    write_pack(path, metadata=metadata(**meta), rows=rows, created=CREATED)
    return path


def _sheet_values(path, sheet):
    wb = openpyxl.load_workbook(path)
    try:
        return [list(row) for row in wb[sheet].iter_rows(values_only=True)]
    finally:
        wb.close()


def test_one_item_pack_validates_from_disk(tmp_path):
    path = _write(tmp_path, [valid_row(1)])
    assert validate_workbook(path) == []


def test_every_value_round_trips_as_str(tmp_path):
    row = valid_row(1)
    path = _write(tmp_path, [row])
    wb = openpyxl.load_workbook(path)
    try:
        for ws in wb.worksheets:
            for cells in ws.iter_rows():
                for cell in cells:
                    if cell.value is not None:
                        assert isinstance(cell.value, str), cell.coordinate
                        assert cell.data_type == "s", cell.coordinate
        evidence = wb[fmt.SHEET_EVIDENCE]
        for column, header in enumerate(fmt.EVIDENCE_HEADERS, start=1):
            assert evidence.cell(row=2, column=column).value == row.get(header), header
    finally:
        wb.close()


def test_raw_xml_has_no_formula_and_only_text_cells(tmp_path):
    rows = [with_values(valid_row(1), {"Description": "=SUM(1,2) in the heading."})]
    path = _write(tmp_path, rows)
    for name, xml in sheet_xml(path).items():
        assert "<f>" not in xml and "<f " not in xml, name
        types = set(re.findall(r'<c [^>]*?t="([^"]+)"', xml))
        assert types <= {"inlineStr", "s", "str"}, (name, types)
        typeless_with_value = re.findall(r"<c (?![^>]*\bt=)[^>]*>\s*<v>", xml)
        assert typeless_with_value == [], name


@pytest.mark.parametrize("value", EDGE_VALUES)
def test_edge_values_survive_as_exact_strings(tmp_path, value):
    row = with_values(
        valid_row(1),
        {
            "Description": value,
            "Collection Conditions": value,
            "Preservation Method": value,
        },
    )
    path = _write(tmp_path, [row])
    values = _sheet_values(path, fmt.SHEET_EVIDENCE)[1]
    for header in ("Description", "Collection Conditions", "Preservation Method"):
        assert values[fmt.EVIDENCE_HEADERS.index(header)] == value


def test_formula_risk_depends_only_on_first_character(tmp_path):
    texts = ("=x", "a=b", "   =SUM(1,2)")
    rows = [with_values(valid_row(n + 1), {"Description": t}) for n, t in enumerate(texts)]
    path = _write(tmp_path, rows)
    values = _sheet_values(path, fmt.SHEET_EVIDENCE)
    column = fmt.EVIDENCE_HEADERS.index("Description")
    assert [values[r][column] for r in (1, 2, 3)] == list(texts)
    assert validate_workbook(path) == []


def _col_widths(path, sheet_part):
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read(sheet_part))
    widths = {}
    for col in root.iter(f"{{{MAIN_NS}}}col"):
        assert col.get("min") == col.get("max"), "one <col> per column"
        widths[int(col.get("min"))] = (col.get("width"), col.get("customWidth"))
    return widths


def test_column_widths_on_disk(tmp_path):
    path = _write(tmp_path, [valid_row(1)])
    widths = _col_widths(path, "xl/worksheets/sheet2.xml")
    assert widths[1] == ("12.83203125", "1")
    assert widths[3] == ("60.83203125", "1")


def test_header_and_width_never_shift(tmp_path):
    path = _write(tmp_path, [valid_row(1)])
    headers = _sheet_values(path, fmt.SHEET_EVIDENCE)[0]
    assert tuple(headers) == fmt.EVIDENCE_HEADERS
    widths = _col_widths(path, "xl/worksheets/sheet2.xml")
    for index, wch in enumerate(fmt.EVIDENCE_WIDTHS_WCH, start=1):
        assert float(widths[index][0]) == wch + fmt.WIDTH_OFFSET, index
    attachment_widths = _col_widths(path, "xl/worksheets/sheet3.xml")
    for index, wch in enumerate(fmt.ATTACHMENT_WIDTHS_WCH, start=1):
        assert float(attachment_widths[index][0]) == wch + fmt.WIDTH_OFFSET, index


def test_sheet_order_and_metadata_layout(tmp_path):
    path = _write(tmp_path, [valid_row(1), valid_row(2)])
    wb = openpyxl.load_workbook(path)
    try:
        assert wb.sheetnames == list(fmt.SHEET_ORDER)
    finally:
        wb.close()
    meta = _sheet_values(path, fmt.SHEET_METADATA)
    assert len(meta) == fmt.METADATA_ROW_COUNT
    assert [r[0] for r in meta[:6]] == list(fmt.METADATA_LABELS)
    assert meta[0][1] == fmt.FORMAT_MARKER
    assert meta[4][1] == "2"
    assert meta[5][1] == "0"
    assert meta[6] == [None, None]
    assert [r[0] for r in meta[7:10]] == list(fmt.METADATA_NOTES)
    attachments = _sheet_values(path, fmt.SHEET_ATTACHMENTS)
    assert attachments == [list(fmt.ATTACHMENT_HEADERS)]


@pytest.mark.parametrize("count", [0, 1, 8])
def test_total_evidence_items_is_plain_decimal(tmp_path, count):
    rows = [valid_row(n + 1) for n in range(count)]
    path = _write(tmp_path, rows)
    meta = _sheet_values(path, fmt.SHEET_METADATA)
    assert meta[4][1] == str(count)
    assert validate_workbook(path) == []
    if count == 0:
        assert _sheet_values(path, fmt.SHEET_EVIDENCE) == [list(fmt.EVIDENCE_HEADERS)]


def test_cells_follow_header_order_not_key_order(tmp_path):
    row = valid_row(1)
    reversed_row = dict(reversed(list(row.items())))
    path = _write(tmp_path, [reversed_row])
    values = _sheet_values(path, fmt.SHEET_EVIDENCE)[1]
    for index, header in enumerate(fmt.EVIDENCE_HEADERS):
        assert values[index] == row.get(header), header


def test_empty_optional_column_writes_no_cell(tmp_path):
    row = with_values(valid_row(1), {"Collection Date": None, "Retention Until": ""})
    path = _write(tmp_path, [row])
    wb = openpyxl.load_workbook(path)
    try:
        ws = wb[fmt.SHEET_EVIDENCE]
        for header in ("Collection Date", "Retention Until", "QA Status"):
            cell = ws.cell(row=2, column=fmt.EVIDENCE_HEADERS.index(header) + 1)
            assert cell.value is None, header
    finally:
        wb.close()


def test_control_characters_removed_and_astral_kept(tmp_path):
    text = "Line one\tcol\nline two\r\x00\x07\x0b\x1fend \U0001f525"
    row = with_values(valid_row(1), {"Description": text})
    path = _write(tmp_path, [row])
    values = _sheet_values(path, fmt.SHEET_EVIDENCE)[1]
    assert values[fmt.EVIDENCE_HEADERS.index("Description")] == (
        "Line one\tcol\nline two\rend \U0001f525"
    )


def test_non_text_value_rejected_before_writing(tmp_path):
    target = tmp_path / "pack.xlsx"
    row = with_values(valid_row(1), {"Collection Conditions": 0})
    with pytest.raises(PackInvalid) as info:
        write_pack(target, metadata=metadata(), rows=[row], created=CREATED)
    assert "CELL_NOT_TEXT" in [f.code for f in info.value.findings]
    assert not target.exists()


def test_failed_disk_validation_leaves_target_untouched(tmp_path, monkeypatch):
    target = tmp_path / "pack.xlsx"
    target.write_bytes(b"previous pack")
    planted = Finding("CELL_FORMULA", fmt.SHEET_EVIDENCE, "C2", "Description", "planted")
    monkeypatch.setattr(writer_module, "validate_workbook", lambda path: [planted])
    with pytest.raises(PackInvalid):
        write_pack(target, metadata=metadata(), rows=[valid_row(1)], created=CREATED)
    assert target.read_bytes() == b"previous pack"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["pack.xlsx"]


def test_validation_reads_the_saved_temp_file(tmp_path, monkeypatch):
    seen = []

    def spy(path):
        seen.append(path.name)
        return validate_module.validate_workbook(path)

    monkeypatch.setattr(writer_module, "validate_workbook", spy)
    _write(tmp_path, [valid_row(1)])
    assert seen == ["pack.xlsx.tmp"]


def test_created_and_modified_properties_are_injected(tmp_path):
    path = _write(tmp_path, [valid_row(1)])
    wb = openpyxl.load_workbook(path)
    try:
        assert wb.properties.created == CREATED
        assert wb.properties.modified == CREATED
    finally:
        wb.close()
