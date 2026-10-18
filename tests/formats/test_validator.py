"""Validator rejection classes (BLD-03, DEF-14): every code fires on its planted defect only.

Row-value defects are checked twice: in memory through ``validate_rows`` and on disk
through ``validate_workbook`` after planting the same text in a valid pack. Structure and
raw-XML defects are planted in a saved workbook.
"""

import json
import zipfile
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest
from epb.formats import probematrix_v2 as fmt
from epb.formats.validate import validate_rows, validate_workbook
from epb.formats.writer import write_pack
from epbtest.packkit import CREATED, CUSTODY_ENTRY, codes, metadata, valid_row, with_values

REPO_ROOT = Path(__file__).resolve().parents[2]
PROTOTYPE = REPO_ROOT / "docs/reference/hazira-2020/prototype-evidence-pack-hazira-2020.xlsx"


_DROP = object()


def _custody(**changes):
    entry = dict(CUSTODY_ENTRY)
    for key, value in changes.items():
        if value is _DROP:
            entry.pop(key)
        else:
            entry[key] = value
    return json.dumps([entry])


# (case id, header, planted text, expected code). Each plants exactly one defect.
TEXT_CASES = [
    ("type-blank", "Type", "   ", "REQUIRED_BLANK"),
    ("description-blank", "Description", " \n ", "REQUIRED_BLANK"),
    ("type-not-exact", "Type", "Witness_Statement", "ENUM_INVALID"),
    ("type-unknown", "Type", "email", "ENUM_INVALID"),
    ("status-trailing-space", "Analysis Status", "pending ", "ENUM_INVALID"),
    ("priority-case", "Priority", "Low", "ENUM_INVALID"),
    ("qa-status-case", "QA Status", "Verified", "ENUM_INVALID"),
    ("qa-status-unknown", "QA Status", "approved", "ENUM_INVALID"),
    ("retention-space", "Retention Category", "legal ", "ENUM_INVALID"),
    ("integrity-lowercase", "Integrity Verified", "yes", "ENUM_INVALID"),
    ("date-unpadded-month", "Collection Date", "2026-9", "DATE_PRECISION"),
    ("date-day-first", "Collection Date", "17/10/2026", "DATE_PRECISION"),
    ("date-with-time", "Collection Date", "2026-09-14T00:00", "DATE_PRECISION"),
    ("custody-not-json", "Chain of Custody (JSON)", "not json", "CUSTODY_JSON"),
    ("custody-object", "Chain of Custody (JSON)", json.dumps(CUSTODY_ENTRY), "CUSTODY_JSON"),
    ("custody-list-of-numbers", "Chain of Custody (JSON)", "[1, 2]", "CUSTODY_JSON"),
    (
        "custody-missing-key",
        "Chain of Custody (JSON)",
        _custody(received_by=_DROP),
        "CUSTODY_ENTRY",
    ),
    (
        "custody-empty-required",
        "Chain of Custody (JSON)",
        _custody(transferred_to=""),
        "CUSTODY_ENTRY",
    ),
    (
        "custody-bad-timestamp",
        "Chain of Custody (JSON)",
        _custody(timestamp="yesterday"),
        "CUSTODY_ENTRY",
    ),
    ("custody-unknown-key", "Chain of Custody (JSON)", _custody(colour="red"), "CUSTODY_ENTRY"),
    (
        "custody-bool-as-string",
        "Chain of Custody (JSON)",
        _custody(verification_confirmed="true"),
        "CUSTODY_ENTRY",
    ),
    (
        "custody-bad-condition",
        "Chain of Custody (JSON)",
        _custody(condition_on_receipt="broken"),
        "CUSTODY_ENTRY",
    ),
    (
        "custody-bad-action",
        "Chain of Custody (JSON)",
        _custody(action_type="moved"),
        "CUSTODY_ENTRY",
    ),
    ("list-whitespace", "Tags", " pump-trip, alarm-log", "LIST_CELL_WHITESPACE"),
    ("list-semicolon", "Linked Equipment", "Pump P-101; Valve V-12", "LIST_SEMICOLON"),
    ("list-no-space", "Tags", "pump-trip,alarm-log", "LIST_SEPARATOR"),
    ("list-two-spaces", "Tags", "pump-trip,  alarm-log", "LIST_SEPARATOR"),
    ("list-empty-fragment", "Related Evidence", "TAG-ACM-1, , TAG-ACM-2", "LIST_EMPTY_FRAGMENT"),
    (
        "list-unbalanced",
        "Linked Equipment",
        "Pump P-101 (main, spare)",
        "LIST_FRAGMENT_UNBALANCED",
    ),
    ("list-duplicate", "Regulatory Requirements", "Rule 7, Rule 7", "LIST_DUPLICATE"),
]


def _ids(cases):
    return [case[0] for case in cases]


@pytest.mark.parametrize(("case", "header", "text", "code"), TEXT_CASES, ids=_ids(TEXT_CASES))
def test_row_defect_in_memory(case, header, text, code):
    findings = validate_rows(metadata(), [with_values(valid_row(1), {header: text})])
    assert codes(findings) == [code]
    assert findings[0].sheet == fmt.SHEET_EVIDENCE
    assert findings[0].column == header
    column = fmt.EVIDENCE_HEADERS.index(header) + 1
    assert findings[0].cell == f"{openpyxl.utils.get_column_letter(column)}2"
    assert header in findings[0].message


def _pack(tmp_path, rows=None, name="pack.xlsx"):
    path = tmp_path / name
    rows = [valid_row(1)] if rows is None else rows
    write_pack(path, metadata=metadata(), rows=rows, created=CREATED)
    return path


def _mutate(path, change):
    """Load the pack with plain openpyxl, apply ``change(wb)`` and save it back."""
    wb = openpyxl.load_workbook(path)
    change(wb)
    wb.save(path)
    wb.close()


def _set(sheet, ref, value):
    def change(wb):
        wb[sheet][ref].value = value

    return change


@pytest.mark.parametrize(("case", "header", "text", "code"), TEXT_CASES, ids=_ids(TEXT_CASES))
def test_row_defect_on_disk(tmp_path, case, header, text, code):
    path = _pack(tmp_path)
    column = openpyxl.utils.get_column_letter(fmt.EVIDENCE_HEADERS.index(header) + 1)
    _mutate(path, _set(fmt.SHEET_EVIDENCE, f"{column}2", text))
    assert codes(validate_workbook(path)) == [code]


@pytest.mark.parametrize(
    ("header", "value"),
    [("Collection Conditions", 0), ("Preservation Method", True), ("Source", datetime(2026, 9, 1))],
)
def test_non_text_value_in_memory(header, value):
    findings = validate_rows(metadata(), [with_values(valid_row(1), {header: value})])
    assert codes(findings) == ["CELL_NOT_TEXT"]


def test_non_text_metadata_value_in_memory():
    findings = validate_rows(metadata(incident_number=2026), [valid_row(1)])
    assert codes(findings) == ["CELL_NOT_TEXT"]
    assert findings[0].sheet == fmt.SHEET_METADATA


def test_metadata_attachments_in_memory():
    findings = validate_rows(metadata(total_attachments="1"), [valid_row(1)])
    assert codes(findings) == ["METADATA_ATTACHMENTS"]


@pytest.mark.parametrize(
    ("length", "expected"),
    [(32766, []), (32767, ["CELL_PROBABLY_TRUNCATED"]), (32768, ["CELL_TOO_LONG"])],
)
def test_length_boundary_in_memory(length, expected):
    row = with_values(valid_row(1), {"Description": "x" * length})
    assert codes(validate_rows(metadata(), [row])) == expected


def _replace_in_part(path, part, old, new):
    with zipfile.ZipFile(path) as archive:
        items = [(info, archive.read(info.filename)) for info in archive.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for info, data in items:
            if info.filename == part:
                text = data.decode("utf-8")
                assert old in text
                data = text.replace(old, new, 1).encode("utf-8")
            archive.writestr(info, data)


@pytest.mark.parametrize(
    ("length", "expected"),
    [(32766, []), (32767, ["CELL_PROBABLY_TRUNCATED"]), (32768, ["CELL_TOO_LONG"])],
)
def test_length_boundary_on_disk(tmp_path, length, expected):
    path = _pack(tmp_path)
    original = valid_row(1)["Description"]
    _replace_in_part(path, "xl/worksheets/sheet2.xml", original, "x" * length)
    assert codes(validate_workbook(path)) == expected


@pytest.mark.parametrize("char", ["￾", "￿", "\ud800", "\udfff"])
def test_bad_characters_rejected(char):
    row = with_values(valid_row(1), {"Description": f"Bad {char} character."})
    assert codes(validate_rows(metadata(), [row])) == ["CELL_BAD_CHAR"]


def test_astral_characters_accepted():
    row = with_values(valid_row(1), {"Description": "Flame \U0001f525 symbol."})
    assert validate_rows(metadata(), [row]) == []


def test_empty_custody_list_accepted():
    row = with_values(valid_row(1), {"Chain of Custody (JSON)": "[]"})
    assert validate_rows(metadata(), [row]) == []


def test_blank_custody_rejected():
    row = with_values(valid_row(1), {"Chain of Custody (JSON)": None})
    assert codes(validate_rows(metadata(), [row])) == ["CUSTODY_JSON"]


def test_custody_entry_message_names_index_and_key():
    text = json.dumps([CUSTODY_ENTRY, {**CUSTODY_ENTRY, "colour": "red"}])
    row = with_values(valid_row(1), {"Chain of Custody (JSON)": text})
    findings = validate_rows(metadata(), [row])
    assert codes(findings) == ["CUSTODY_ENTRY"]
    assert "entry 2" in findings[0].message
    assert "colour" in findings[0].message


def test_all_optional_custody_keys_accepted():
    entry = {
        **CUSTODY_ENTRY,
        "reason": "copy for analysis",
        "verification_confirmed": False,
        "location": "case folder",
        "notes": "synthetic",
    }
    row = with_values(valid_row(1), {"Chain of Custody (JSON)": json.dumps([entry])})
    assert validate_rows(metadata(), [row]) == []


@pytest.mark.parametrize("value", sorted(fmt.QA_STATUSES))
def test_every_qa_status_accepted(value):
    row = with_values(valid_row(1), {"QA Status": value})
    assert validate_rows(metadata(), [row]) == []


@pytest.mark.parametrize("value", ["2026", "2026-09", "2026-09-14"])
def test_collection_date_precisions_accepted(value):
    row = with_values(valid_row(1), {"Collection Date": value})
    assert validate_rows(metadata(), [row]) == []


def test_well_formed_lists_accepted():
    row = with_values(
        valid_row(1),
        {
            "Tags": "pump-trip, alarm-log, shift-handover",
            "Linked Equipment": "Pump P-101 (main), Valve V-12",
        },
    )
    assert validate_rows(metadata(), [row]) == []


def test_one_finding_per_cell_and_code():
    row = with_values(valid_row(1), {"Linked Equipment": "A (x, y); B (z, w)"})
    found = Counter(codes(validate_rows(metadata(), [row])))
    assert found == Counter({"LIST_SEMICOLON": 1, "LIST_FRAGMENT_UNBALANCED": 1})


# --- Structure defects, planted in a saved workbook -------------------------------------


def _swap_metadata_rows(wb):
    ws = wb[fmt.SHEET_METADATA]
    for column in ("A", "B"):
        upper, lower = ws[f"{column}2"].value, ws[f"{column}3"].value
        ws[f"{column}2"].value, ws[f"{column}3"].value = lower, upper


def _swap_headers(wb):
    ws = wb[fmt.SHEET_EVIDENCE]
    by = fmt.EVIDENCE_HEADERS.index("QA Reviewed By") + 1
    at = fmt.EVIDENCE_HEADERS.index("QA Reviewed At") + 1
    a, b = ws.cell(row=1, column=by), ws.cell(row=1, column=at)
    a.value, b.value = b.value, a.value


def _reorder_sheets(wb):
    wb.move_sheet(fmt.SHEET_ATTACHMENTS, offset=-1)


def _drop_attachments(wb):
    del wb[fmt.SHEET_ATTACHMENTS]


def _extra_sheet(wb):
    wb.create_sheet("Notes")


STRUCTURE_CASES = [
    ("sheets-reordered", _reorder_sheets, "SHEETS_INVALID"),
    ("sheets-missing", _drop_attachments, "SHEETS_INVALID"),
    ("sheets-extra", _extra_sheet, "SHEETS_INVALID"),
    ("metadata-swapped-rows", _swap_metadata_rows, "METADATA_LAYOUT"),
    ("metadata-row7-filled", _set(fmt.SHEET_METADATA, "A7", "extra"), "METADATA_LAYOUT"),
    ("metadata-row11", _set(fmt.SHEET_METADATA, "A11", "extra"), "METADATA_LAYOUT"),
    ("metadata-note-changed", _set(fmt.SHEET_METADATA, "A9", "Changed."), "METADATA_LAYOUT"),
    ("marker-v1", _set(fmt.SHEET_METADATA, "B1", "EVIDENCE_PACK_V1"), "FORMAT_MARKER"),
    ("total-wrong", _set(fmt.SHEET_METADATA, "B5", "2"), "METADATA_TOTAL_MISMATCH"),
    ("total-padded", _set(fmt.SHEET_METADATA, "B5", "01"), "METADATA_TOTAL_MISMATCH"),
    ("total-decimal", _set(fmt.SHEET_METADATA, "B5", "1.0"), "METADATA_TOTAL_MISMATCH"),
    ("attachments-total", _set(fmt.SHEET_METADATA, "B6", "1"), "METADATA_ATTACHMENTS"),
    ("headers-swapped", _swap_headers, "HEADERS_INVALID"),
    ("attachments-row", _set(fmt.SHEET_ATTACHMENTS, "A2", "EVD-ACM-001"), "ATTACHMENTS_INVALID"),
    (
        "attachments-header",
        _set(fmt.SHEET_ATTACHMENTS, "C1", "Filename"),
        "ATTACHMENTS_INVALID",
    ),
    ("formula", _set(fmt.SHEET_EVIDENCE, "S2", "=1+1"), "CELL_FORMULA"),
    ("number", _set(fmt.SHEET_EVIDENCE, "T2", 5), "CELL_NOT_TEXT"),
    ("boolean", _set(fmt.SHEET_EVIDENCE, "U2", True), "CELL_NOT_TEXT"),
]


@pytest.mark.parametrize(
    ("case", "change", "code"), STRUCTURE_CASES, ids=[c[0] for c in STRUCTURE_CASES]
)
def test_structure_defect_on_disk(tmp_path, case, change, code):
    path = _pack(tmp_path)
    _mutate(path, change)
    assert set(codes(validate_workbook(path))) == {code}


def test_plain_openpyxl_number_formula_boolean(tmp_path):
    path = _pack(tmp_path)

    def change(wb):
        ws = wb[fmt.SHEET_EVIDENCE]
        ws["S2"].value = 42
        ws["T2"].value = "=SUM(1,2)"
        ws["U2"].value = False

    _mutate(path, change)
    found = Counter((f.code, f.cell) for f in validate_workbook(path))
    assert found == Counter({("CELL_NOT_TEXT", "S2"): 1, ("CELL_FORMULA", "T2"): 1,
                             ("CELL_NOT_TEXT", "U2"): 1})  # fmt: skip


def test_str_cell_without_formula_accepted(tmp_path):
    path = _pack(tmp_path)
    original = valid_row(1)["Description"]
    _replace_in_part(
        path,
        "xl/worksheets/sheet2.xml",
        f'<c r="C2" s="1" t="inlineStr"><is><t>{original}</t></is></c>',
        f'<c r="C2" s="1" t="str"><v>{original}</v></c>',
    )
    assert validate_workbook(path) == []


def test_typeless_cell_with_value_rejected(tmp_path):
    path = _pack(tmp_path)
    original = valid_row(1)["Description"]
    _replace_in_part(
        path,
        "xl/worksheets/sheet2.xml",
        f'<c r="C2" s="1" t="inlineStr"><is><t>{original}</t></is></c>',
        '<c r="C2" s="1"><v>7</v></c>',
    )
    assert "CELL_NOT_TEXT" in codes(validate_workbook(path))


def test_not_a_workbook(tmp_path):
    path = tmp_path / "pack.xlsx"
    path.write_bytes(b"this is not a zip file")
    assert codes(validate_workbook(path)) == ["SHEETS_INVALID"]


def test_broken_sheet_xml(tmp_path):
    path = _pack(tmp_path)
    _replace_in_part(path, "xl/worksheets/sheet2.xml", "</sheetData>", "</sheetDatum>")
    assert "SHEETS_INVALID" in codes(validate_workbook(path))


def test_zero_item_pack_validates(tmp_path):
    assert validate_workbook(_pack(tmp_path, rows=[])) == []


def test_findings_sorted_and_stable(tmp_path):
    rows = [
        with_values(valid_row(1), {"Priority": "Low", "Tags": "a,b"}),
        with_values(valid_row(2), {"Type": "", "QA Status": "x"}),
    ]
    path = _pack(tmp_path)
    _mutate(path, _set(fmt.SHEET_METADATA, "B1", "EVIDENCE_PACK_V1"))
    first = validate_workbook(path)
    assert first == validate_workbook(path)
    in_memory = validate_rows(metadata(), rows)
    assert in_memory == validate_rows(metadata(), list(rows))
    assert [(f.cell, f.code) for f in in_memory] == [
        ("D2", "ENUM_INVALID"),
        ("F2", "LIST_SEPARATOR"),
        ("B3", "REQUIRED_BLANK"),
        ("Y3", "ENUM_INVALID"),
    ]


# --- The prototype reference pack (codes and counts only; no cell text) ----------------


def test_prototype_pack_fails_on_list_fragments():
    findings = validate_workbook(PROTOTYPE)
    found = Counter((f.code, f.column) for f in findings)
    assert "SHEETS_INVALID" in codes(findings)
    assert found[("LIST_FRAGMENT_UNBALANCED", "Linked Equipment")] == 24
    assert found[("LIST_SEMICOLON", "Linked Equipment")] == 24
    assert found[("LIST_SEMICOLON", "Linked Individuals")] == 1
